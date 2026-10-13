"""
Segunda tentativa de geocodificação para endereços não encontrados.

Para cada endereço que ficou sem coordenada no cache_geo.json, tenta:
  1. Expandir a abreviação do logradouro (AVE → AVENIDA, etc.)
  2. Remover o número (só rua + bairro + cidade)
  3. Só rua + cidade (sem bairro)

Atualiza o cache_geo.json e regera o bares_com_geohash.xlsx ao final.
"""

import json
import time
import openpyxl
from geopy.geocoders import Nominatim
from geopy.exc import GeocoderTimedOut, GeocoderServiceError
from geohash import encode as geohash_encode

ARQUIVO_ENTRADA  = "bares.xlsx"
ARQUIVO_SAIDA    = "bares_com_geohash.xlsx"
ARQUIVO_CACHE    = "cache_geo.json"
PRECISAO_GEOHASH = 8
PAUSA_SEGUNDOS   = 1.1

# Abreviações mais comuns em endereços de BH
ABREVIACOES = {
    "AVE":  "AVENIDA",
    "AL":   "ALAMEDA",
    "PC":   "PRAÇA",
    "PCA":  "PRAÇA",
    "EST":  "ESTRADA",
    "ROD":  "RODOVIA",
    "TV":   "TRAVESSA",
    "BEC":  "BECO",
    "LGO":  "LARGO",
    "VL":   "VILA",
    "CJ":   "CONJUNTO",
    "R":    "RUA",
}


def expandir_abreviacao(desc_log):
    return ABREVIACOES.get(desc_log.upper(), desc_log)


def carregar_cache():
    try:
        with open(ARQUIVO_CACHE, "r", encoding="utf-8") as f:
            return json.load(f)
    except FileNotFoundError:
        return {}


def salvar_cache(cache):
    with open(ARQUIVO_CACHE, "w", encoding="utf-8") as f:
        json.dump(cache, f, ensure_ascii=False, indent=2)


def consultar_nominatim(geolocator, endereco):
    try:
        resultado = geolocator.geocode(endereco, timeout=10)
        if resultado:
            return resultado.latitude, resultado.longitude
        return None, None
    except (GeocoderTimedOut, GeocoderServiceError) as e:
        print(f"  ERRO: {e}")
        return None, None


# -----------------------------------------------------------------------
# Ler planilha e reconstruir a lista de endereços originais
# -----------------------------------------------------------------------

print("Lendo planilha...")
wb = openpyxl.load_workbook(ARQUIVO_ENTRADA, read_only=True, data_only=True)
ws = wb.active
todas_as_linhas = list(ws.iter_rows(values_only=True))
cabecalho = todas_as_linhas[0]
dados = todas_as_linhas[1:]
wb.close()

# -----------------------------------------------------------------------
# Montar endereços alternativos para os que falharam
# -----------------------------------------------------------------------

cache = carregar_cache()

# Identifica quais endereços únicos falharam
falhos = {end for end, val in cache.items() if val["lat"] is None}
print(f"Endereços sem coordenada no cache: {len(falhos)}")

geolocator = Nominatim(user_agent="tp_alg2_bh_geohash_v2")
recuperados = 0

for row in dados:
    desc_log = str(row[3] or "").strip()
    nome_log = str(row[4] or "").strip()
    numero   = str(row[5] or "").strip()
    bairro   = str(row[7] or "").strip()

    # Endereço original (mesmo formato do geocodificar.py)
    end_original = f"{desc_log} {nome_log}, {numero}, {bairro}, Belo Horizonte, MG, Brasil"

    if end_original not in falhos:
        continue   # já tem coordenada ou não está no cache, pula

    # Remove da lista para não processar a mesma rua duas vezes
    falhos.discard(end_original)

    desc_expandido = expandir_abreviacao(desc_log)

    # Formatos alternativos, do mais específico ao mais genérico
    tentativas = [
        f"{desc_expandido} {nome_log}, {numero}, {bairro}, Belo Horizonte, MG, Brasil",
        f"{desc_expandido} {nome_log}, {bairro}, Belo Horizonte, MG, Brasil",
        f"{desc_expandido} {nome_log}, Belo Horizonte, MG, Brasil",
        f"{nome_log}, {bairro}, Belo Horizonte, MG, Brasil",
    ]

    lat, lon = None, None
    tentativa_usada = None

    for tentativa in tentativas:
        lat, lon = consultar_nominatim(geolocator, tentativa)
        time.sleep(PAUSA_SEGUNDOS)
        if lat is not None:
            tentativa_usada = tentativa
            break

    # Atualiza o cache com o resultado (encontrado ou não)
    cache[end_original] = {"lat": lat, "lon": lon}
    salvar_cache(cache)

    if lat is not None:
        recuperados += 1
        print(f"  RECUPERADO via '{tentativa_usada[:50]}' → lat={lat:.4f}, lon={lon:.4f}")
    else:
        print(f"  Ainda não encontrado: {end_original[:60]}")

print(f"\nEndereços recuperados nesta rodada: {recuperados}")

# -----------------------------------------------------------------------
# Regerar o arquivo de saída com o cache atualizado
# -----------------------------------------------------------------------

print("\nRegerando bares_com_geohash.xlsx...")

wb_saida = openpyxl.Workbook()
ws_saida = wb_saida.active
ws_saida.title = "Resultado"
ws_saida.append(list(cabecalho) + ["LATITUDE", "LONGITUDE", "GEOHASH", "ENDERECO_CONSULTADO"])

sem_coordenada = 0

for row in dados:
    desc_log = str(row[3] or "").strip()
    nome_log = str(row[4] or "").strip()
    numero   = str(row[5] or "").strip()
    bairro   = str(row[7] or "").strip()
    end_original = f"{desc_log} {nome_log}, {numero}, {bairro}, Belo Horizonte, MG, Brasil"

    resultado = cache.get(end_original, {})
    lat = resultado.get("lat")
    lon = resultado.get("lon")

    if lat is not None and lon is not None:
        geohash = geohash_encode(lat, lon, PRECISAO_GEOHASH)
    else:
        geohash = ""
        sem_coordenada += 1

    ws_saida.append(list(row) + [lat, lon, geohash, end_original])

wb_saida.save(ARQUIVO_SAIDA)
print(f"Arquivo salvo: {ARQUIVO_SAIDA}")
print(f"  Com Geohash: {len(dados) - sem_coordenada}")
print(f"  Sem coordenada: {sem_coordenada}")
